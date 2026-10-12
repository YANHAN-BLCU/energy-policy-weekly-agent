"""生成国家级周报 Excel：首表 Police_<yymmdd>_Test + Issues 附表。

输入：运行目录下的 analysis.json（宿主 AI 生成的摘要/分类，字段：
       policy_id / index / supported_industry / policy_tags / issuing_agencies /
       policy_abstract / impact_analysis）
      以及 run 目录下 collect_static 产出的 analysis_request.json / failed_details.json /
      review_needed.json / run.json

用法：
    python3 scripts/build_workbook.py <run_dir>
"""
import json
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[1]

HEADERS = ["Policy_ID", "Policy_Name", "Policy_Number", "Issue_Date", "Issuing_Agency",
           "Issuing_Agency_Name", "Supported_Industry", "Policy_Tags", "Policy_Abstract",
           "Impact_Analysis", "Full_Text_Link", "Source_List_URL", "Attachment_Links",
           "Collected_At", "Collector_Method", "Collection_Status", "AI_Status", "Evidence_Path"]
ISSUE_HEADERS = ["Issue_Type", "Agency", "Policy_Title", "Issue", "Possible_Omission",
                 "Detail_URL", "Source_List_URL"]


def _url_cell(ws, row, col, value):
    cell = ws.cell(row, col, value)
    if value:
        parsed = urlsplit(str(value))
        if parsed.scheme in {"http", "https"} and parsed.netloc:
            cell.hyperlink = str(value)
            cell.font = Font(color="0563C1", underline="single")


def _code_names(issuing_agencies, dictionary):
    by_code = {a["code"]: a for a in dictionary.get("agencies", []) + dictionary.get("issuer_units", [])}
    codes, names = [], []
    for item in issuing_agencies or []:
        code = item.get("code", "") if isinstance(item, dict) else str(item)
        if not code:
            continue
        codes.append(code)
        names.append(by_code.get(code, {}).get("name", ""))
    return " / ".join(codes), ", ".join(names)


def build(run_dir):
    run_dir = Path(run_dir).resolve()
    run = json.loads((run_dir / "run.json").read_text(encoding="utf-8"))
    request = json.loads((run_dir / "analysis_request.json").read_text(encoding="utf-8"))
    dictionary = request.get("dictionary", {})
    analysis_path = run_dir / "analysis.json"
    analysis = json.loads(analysis_path.read_text(encoding="utf-8")) if analysis_path.exists() else {}

    by_idx = {}
    for item in analysis.get("records", []):
        by_idx[item.get("index", item.get("policy_id"))] = item

    records = request.get("records", [])
    records.sort(key=lambda r: (r.get("issue_date", ""), r.get("title", "")), reverse=True)

    day = run.get("window_end", datetime.now().date().isoformat())
    province = (run.get("province") or "").strip()
    suffix = province if province else "Regional"
    sheet_name = "Police_" + day.replace("-", "")[2:8] + "_Test"
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(color="FFFFFF", bold=True)
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    collected_at = run.get("generated_at", "")
    for idx, r in enumerate(records):
        a = by_idx.get(idx, by_idx.get(r.get("title")))
        code_str, name_str = _code_names((a or {}).get("issuing_agencies"), dictionary)
        policy_id = (a or {}).get("policy_id", f"P-{idx+1:04d}")
        # 附件链接：多个用 " ; " 分隔
        atts = r.get("attachments") or []
        att_links = " ; ".join(x.get("url", "") if isinstance(x, dict) else str(x) for x in atts)
        ws.append([
            policy_id, r.get("title", ""), r.get("policy_number", ""), r.get("issue_date", ""),
            code_str, name_str,
            (a or {}).get("supported_industry", ""), (a or {}).get("policy_tags", ""),
            (a or {}).get("policy_abstract", ""), (a or {}).get("impact_analysis", ""),
            r.get("detail_url", ""), r.get("list_url", ""), att_links, collected_at,
            "static_html", "Collected", "Host AI analyzed" if a else "Not analyzed",
            r.get("evidence_path", ""),
        ])
        _url_cell(ws, ws.max_row, 11, r.get("detail_url", ""))
        _url_cell(ws, ws.max_row, 12, r.get("list_url", ""))
        # 附件列：单个附件时做成可点超链接
        if len(atts) == 1:
            _url_cell(ws, ws.max_row, 13, atts[0].get("url", "") if isinstance(atts[0], dict) else str(atts[0]))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [20, 44, 22, 14, 20, 46, 26, 34, 60, 60, 50, 50, 50, 22, 16, 16, 18, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    # Issues 附表
    iws = wb.create_sheet("Issues")
    iws.append(ISSUE_HEADERS)
    for cell in iws[1]:
        cell.fill = fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    issues = []
    for r in json.loads((run_dir / "failed_details.json").read_text(encoding="utf-8")).get("records", []):
        issues.append(["BODY_FAILED", r.get("agency", ""), r.get("title", ""),
                       r.get("error", "正文网页文字不足或访问失败"),
                       "未抓取附件；该政策不进入摘要表。", r.get("detail_url", ""), r.get("list_url", "")])
    for r in json.loads((run_dir / "review_needed.json").read_text(encoding="utf-8")).get("items", []):
        issues.append([r.get("kind", "REVIEW"), r.get("agency", ""), r.get("title", ""),
                       r.get("reason", ""), "需宿主 AI 搜索复查后确认。",
                       r.get("detail_url", ""), r.get("list_url", "")])
    if not issues:
        issues.append(["NO_RESULT", "", "", "本次没有通过正文验收的政策记录。",
                       "不代表所有国家级来源本周无政策。", "", ""])
    for issue in issues:
        iws.append(issue)
        if issue[5]:
            _url_cell(iws, iws.max_row, 6, issue[5])
        if issue[6]:
            _url_cell(iws, iws.max_row, 7, issue[6])
    for i, w in enumerate([16, 16, 44, 60, 46, 48, 48], 1):
        iws.column_dimensions[get_column_letter(i)].width = w
    for row in iws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 32

    out = run_dir / f"Police_{day.replace('-', '')[2:]}_Test_{suffix}.xlsx"
    wb.save(out)
    print(json.dumps({"workbook": str(out), "rows": len(records), "issues": len(issues)}, ensure_ascii=False))
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("usage: build_workbook.py <run_dir>", file=sys.stderr)
        raise SystemExit(2)
    build(sys.argv[1])
