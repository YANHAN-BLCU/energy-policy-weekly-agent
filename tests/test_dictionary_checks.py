import copy
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import unittest
import subprocess
import sys
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook, load_workbook

from policy_agent.mi_checks import check_dictionary_assets, check_reference, check_workbook, render_dictionary
from policy_agent.paths import PACKAGE_ROOT
from policy_agent.taxonomy import dictionary
from policy_agent.export import export_workbook


class DictionaryChecksTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix=".test-dictionary-", dir=PACKAGE_ROOT))
        self.addCleanup(lambda: shutil.rmtree(self.root, ignore_errors=True))

    def copy_assets(self):
        for relative in ("policy_agent/data/mi_dictionary.json", "docs/MI_DICTIONARY.md", "PACKAGE_MANIFEST.json"):
            path = self.root / relative
            path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(PACKAGE_ROOT / relative, path)

    def workbook(self):
        row = {"body_valid": True, "body_text": "示例能源机构关于节能政策的通知", "review_status": "Collected",
               "supported_industry": "Example Energy", "policy_tags": "Example Efficiency",
               "issuing_agencies": [{"code": "DEMO", "evidence_quote": "示例能源机构"}],
               "detail_url": "https://example.com/policy", "list_url": "https://example.com/list"}
        path = self.root / "Policy_Weekly.xlsx"
        export_workbook([row], [], path, {"collected_at": "2026-10-09"})
        return path

    def reference(self):
        data = copy.deepcopy(dictionary())
        wb = Workbook()
        ws = wb.active
        ws.title = "Agency abb."
        ws.append(["Code", "English", "Chinese"])
        for agency in data["agencies"]:
            ws.append([agency["code"], agency["name"], agency["chinese_name"]])
        tags = wb.create_sheet("Policy Tags")
        for pair in data["policy_tags"]:
            cell = tags[pair["reference_cell"]]
            cell.value = pair["tag"]
            tags.cell(1, cell.column).value = pair["category"]
        path = self.root / "reference.xlsx"
        wb.save(path)
        raw_image = b"fixture-for-unchanged-image-bytes"
        with ZipFile(path, "a", ZIP_DEFLATED) as archive:
            archive.writestr("xl/media/image1.png", raw_image)
        data["issuing_agency_images"] = [{"part": "xl/media/image1.png", "sha256": hashlib.sha256(raw_image).hexdigest()}]
        return path, data

    def test_relocated_assets_match_manifest(self):
        self.copy_assets()
        result = check_dictionary_assets(self.root, verify_manifest=True)
        self.assertTrue(result["manifest_checked"])
        self.assertEqual(result["tag_pairs"], 1)

    def test_checker_json_remains_utf8_in_isolated_python(self):
        result = subprocess.run([sys.executable, "-I", "-B", str(PACKAGE_ROOT / "scripts/check_mi_dictionary.py")],
                                capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8"))
        report = json.loads(result.stdout.decode("utf-8"))
        self.assertTrue(report["manifest_checked"])
        self.assertFalse(report["reference_checked"])
        self.assertIn("原表未提供", report["reference_status"])

    def test_document_drift_is_detected(self):
        self.copy_assets()
        path = self.root / "docs/MI_DICTIONARY.md"
        path.write_text(path.read_text(encoding="utf-8").replace("Example Energy Agency", "Wrong Agency", 1), encoding="utf-8")
        with self.assertRaises(ValueError):
            check_dictionary_assets(self.root)

    def test_simultaneous_document_and_json_change_cannot_bypass_manifest(self):
        self.copy_assets()
        path = self.root / "policy_agent/data/mi_dictionary.json"
        data = json.loads(path.read_text(encoding="utf-8"))
        data["agencies"][0]["name"] = "Unexpected changed agency"
        for example in data["issuing_agency_examples"]:
            example["english_name"] = example["english_name"].replace("Example Energy Agency", "Unexpected changed agency")
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        (self.root / "docs/MI_DICTIONARY.md").write_text(render_dictionary(data), encoding="utf-8")
        self.assertTrue(check_dictionary_assets(self.root)["assets_consistent"])
        with self.assertRaises(ValueError):
            check_dictionary_assets(self.root, verify_manifest=True)

    def test_reference_compares_values_not_only_record_counts(self):
        path, data = self.reference()
        with patch("policy_agent.mi_checks.dictionary", return_value=data):
            result = check_reference(path)
            self.assertTrue(result["reference_dictionary_consistent"])
            self.assertFalse(result["reference_file_hash_matches"])
            altered = load_workbook(path)
            altered["Agency abb."]["B2"] = "Incorrect translation"
            altered.save(path)
            with self.assertRaisesRegex(ValueError, "Agency abb"):
                check_reference(path)

    def test_reference_detects_wrong_tag_category(self):
        path, data = self.reference()
        altered = load_workbook(path)
        altered["Policy Tags"]["B2"] = "Solar"
        altered.save(path)
        with patch("policy_agent.mi_checks.dictionary", return_value=data), self.assertRaisesRegex(ValueError, "Policy Tags"):
            check_reference(path)

    def test_reference_image_change_requires_review(self):
        path, data = self.reference()
        data["issuing_agency_images"][0]["sha256"] = "0" * 64
        with patch("policy_agent.mi_checks.dictionary", return_value=data), self.assertRaisesRegex(ValueError, "图片"):
            check_reference(path)

    def test_saved_workbook_drift_is_blocked(self):
        path = self.workbook()
        self.assertEqual(check_workbook(path)["policies_checked"], 1)
        for sheet, cell, value in [("Weekly_Policies", "F2", "Wrong translation"),
                                   ("Weekly_Policies", "H2", "Solar"),
                                   ("Tag_Dictionary", "B2", "Wrong tag")]:
            path = self.workbook()
            wb = load_workbook(path)
            wb[sheet][cell] = value
            wb.save(path)
            with self.assertRaises(ValueError):
                check_workbook(path)

    def test_plain_text_and_wrong_hyperlink_targets_are_blocked(self):
        for target in (None, "https://example.com/wrong"):
            path = self.workbook()
            wb = load_workbook(path)
            wb["Weekly_Policies"]["K2"].hyperlink = target
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "超链接"):
                check_workbook(path)
