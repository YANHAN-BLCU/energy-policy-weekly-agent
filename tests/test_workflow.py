import json
import hashlib
import shutil
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from policy_agent.analysis import apply_analysis
from policy_agent.cli import main, read_json, write_json
from policy_agent.export import policy_id, export_workbook
from policy_agent.host_delivery import validate_receipt
from policy_agent.paths import PACKAGE_ROOT, local_path
from policy_agent.setup_wizard import run


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix=".test-runtime-", dir=PACKAGE_ROOT))
        self.addCleanup(lambda: shutil.rmtree(self.root, ignore_errors=True))
        self.installation = self.root / "installation.json"
        run(self.installation, "reader@example.com", "MON", "06:20")
        self.sources = self.root / "sources.json"
        write_json(self.sources, {"sources": [{"code": "TEST", "name": "Test Agency", "website": "https://example.com", "list_url": "https://example.com"}]})
        self.out = self.root / "runs"

    def row(self, key, issued):
        body = "示例能源机构：节能改造项目应符合原文列明的申报要求。"
        evidence = self.root / (key + ".html")
        evidence.write_text(body, encoding="utf-8")
        return {"detail_url": f"https://example.com/{key}.html", "issue_date": issued,
                "title": "节能政策通知", "policy_number": "文号〔2026〕1号",
                "body_text": body, "review_status": "Collected", "body_valid": True,
                "evidence_path": str(evidence), "body_sha256": hashlib.sha256(body.encode("utf-8")).hexdigest()}

    def collect(self, rows):
        with patch("policy_agent.cli.WeeklyCollector.collect", return_value=(rows, [])):
            code = main(["--installation", str(self.installation), "--config", str(self.sources), "--out-dir", str(self.out)])
        self.assertEqual(code, 0)
        return next(self.out.glob("*/run.json")).parent

    def finalize(self, run_dir):
        data = read_json(run_dir / "run.json")
        analysis = {"run_id": data["meta"]["run_id"], "records": [
            {"policy_id": row["policy_id"], "supported_industry": "Example Energy", "policy_tags": "Example Efficiency",
             "issuing_agencies": [{"code": "DEMO", "evidence_quote": "示例能源机构"}],
             "abstract": "政策列明节能项目申报要求。", "impact_analysis": "待对照企业适用范围核查。"}
            for row in data["records"]]}
        write_json(run_dir / "analysis.json", analysis)
        self.assertEqual(main(["--stage", "finalize", "--run-dir", str(run_dir)]), 0)

    def test_setup_has_only_business_values(self):
        data = read_json(self.installation)
        self.assertEqual(set(data), {"initialized", "delivery_email", "weekday", "trigger_time", "scope"})
        with self.assertRaises(ValueError):
            run(self.installation, "bad-email", "MON", "06:20")

    def test_saved_table_mismatch_blocks_delivery_request(self):
        run_dir = self.collect([self.row("new", date.today().isoformat())])
        data = read_json(run_dir / "run.json")
        write_json(run_dir / "analysis.json", {"run_id": data["meta"]["run_id"], "records": [
            {"policy_id": data["records"][0]["policy_id"], "supported_industry": "Example Energy",
             "policy_tags": "Example Efficiency", "abstract": "摘要", "impact_analysis": "影响分析",
             "issuing_agencies": [{"code": "DEMO", "evidence_quote": "示例能源机构"}]}]})
        with patch("policy_agent.cli.check_workbook", side_effect=ValueError("超链接不一致")):
            self.assertEqual(main(["--stage", "finalize", "--run-dir", str(run_dir)]), 2)
        self.assertFalse((run_dir / "delivery_request.json").exists())
        self.assertEqual(read_json(run_dir / "run.json")["meta"]["phase"], "awaiting_analysis")

    def test_external_installation_path_is_rejected(self):
        external = PACKAGE_ROOT.parent / "external-installation.json"
        with self.assertRaises(ValueError):
            run(external, "reader@example.com", "MON", "06:20")

    def test_external_output_and_metadata_state_are_rejected(self):
        external = PACKAGE_ROOT.parent / "outside-policy-workspace"
        self.assertEqual(main(["--installation", str(self.installation), "--config", str(self.sources),
                               "--out-dir", str(external)]), 2)
        with self.assertRaises(ValueError):
            write_json(external / "config.json", {})
        with self.assertRaises(ValueError):
            local_path("../outside-policy-workspace/file.json", "test")

    def test_tampered_receipt_paths_cannot_write_external_state(self):
        run_dir = self.collect([self.row("new", date.today().isoformat())])
        self.finalize(run_dir)
        request = read_json(run_dir / "delivery_request.json")
        receipt_file = run_dir / "receipt.json"
        write_json(receipt_file, {**request, "status": "sent", "transport": "test", "message_id": "simulated-ack"})
        data = read_json(run_dir / "run.json")
        data["meta"]["state_file"] = str(PACKAGE_ROOT.parent / "outside-state.json")
        write_json(run_dir / "run.json", data)
        self.assertEqual(main(["--stage", "receipt", "--run-dir", str(run_dir), "--receipt-file", str(receipt_file)]), 2)

    def test_analysis_cannot_overwrite_collected_facts(self):
        row = self.row("new", date.today().isoformat())
        item = {"policy_id": policy_id(row), "supported_industry": "Example Energy", "policy_tags": "Example Efficiency",
                "issuing_agencies": [{"code": "DEMO", "evidence_quote": "示例能源机构"}],
                "abstract": "摘要", "impact_analysis": "分析", "issue_date": "2099-01-01"}
        enriched = apply_analysis([row], {"run_id": "test", "records": [item]}, "test")
        self.assertEqual(enriched[0]["issue_date"], row["issue_date"])
        self.assertEqual(enriched[0]["policy_number"], row["policy_number"])
        with self.assertRaises(ValueError):
            apply_analysis([row], {"run_id": "test", "records": []}, "test")

    def test_weekly_window_and_missing_date_review(self):
        run_dir = self.collect([self.row("new", date.today().isoformat()),
                                self.row("old", (date.today() - timedelta(days=30)).isoformat()),
                                self.row("unknown", "")])
        data = read_json(run_dir / "run.json")
        self.assertEqual(len(data["records"]), 1)
        self.assertEqual(data["meta"]["outside_window"], 1)
        self.assertEqual(read_json(run_dir / "date_review.json")["records"][0]["review_status"], "Needs review")
        self.assertEqual(len(read_json(run_dir / "analysis_request.json")["records"]), 1)
        self.assertFalse((self.out / "state" / "seen_urls.json").exists())

    def test_failed_body_never_enters_analysis_or_success_workbook(self):
        failed = self.row("column", date.today().isoformat())
        failed.update({"body_valid": False, "body_text": "", "review_status": "Failed", "error": "COLUMN_PAGE"})
        run_dir = self.collect([failed])
        self.assertEqual(read_json(run_dir / "analysis_request.json")["records"], [])
        self.assertEqual(len(read_json(run_dir / "failed_details.json")["records"]), 1)
        with self.assertRaises(ValueError):
            apply_analysis([failed], {"run_id": "test", "records": []}, "test")
        with self.assertRaises(ValueError):
            export_workbook([failed], [], self.root / "invalid.xlsx", {"collected_at": "test"})
        self.finalize(run_dir)
        wb = load_workbook(run_dir / "Policy_Weekly.xlsx", read_only=True)
        self.assertEqual(wb["Weekly_Policies"].max_row, 1)
        wb.close()
        request = read_json(run_dir / "delivery_request.json")
        self.assertIn("采集失败不能", request["body"])
        self.assertNotIn("本周无政策", request["body"].replace("采集失败不能据此判断该站本周无政策", ""))

    def test_delivery_failure_keeps_records_and_sent_commits_only_complete(self):
        run_dir = self.collect([self.row("new", date.today().isoformat()), self.row("unknown", "")])
        self.finalize(run_dir)
        workbook = run_dir / "Policy_Weekly.xlsx"
        original_bytes = workbook.read_bytes()
        wb = load_workbook(workbook, read_only=True)
        self.assertEqual(wb.sheetnames, ["Weekly_Policies", "Run_Summary", "Source_Status", "Tag_Dictionary"])
        wb.close()
        request = read_json(run_dir / "delivery_request.json")
        receipt = {**request, "status": "failed", "transport": "test-connector", "message_id": "", "error": "simulated failure"}
        receipt_file = run_dir / "receipt.json"
        write_json(receipt_file, receipt)
        command = ["--stage", "receipt", "--run-dir", str(run_dir), "--receipt-file", str(receipt_file)]
        self.assertEqual(main(command), 0)
        self.assertFalse((self.out / "state" / "seen_urls.json").exists())
        receipt.update({"status": "sent", "message_id": "simulated-ack-1", "error": ""})
        write_json(receipt_file, receipt)
        self.assertEqual(main(command), 0)
        self.assertEqual(main(command), 0)
        self.assertEqual(read_json(self.out / "state" / "seen_urls.json")["seen_urls"], ["https://example.com/new.html"])
        self.assertEqual(workbook.read_bytes(), original_bytes)
        self.assertEqual(read_json(run_dir / "run.json")["meta"]["phase"], "delivered")
        self.assertEqual(main(["--stage", "finalize", "--run-dir", str(run_dir)]), 2)

    def test_missing_or_wrong_ack_cannot_confirm_delivery(self):
        run_dir = self.collect([self.row("new", date.today().isoformat())])
        self.finalize(run_dir)
        request = read_json(run_dir / "delivery_request.json")
        with self.assertRaises(ValueError):
            validate_receipt(request, {**request, "status": "sent"})
        receipt = {**request, "status": "sent", "transport": "test", "message_id": "simulated-ack"}
        with self.assertRaises(ValueError):
            validate_receipt(request, {**receipt, "recipient": "someone-else@example.com"})
        Path(request["attachment"]).write_bytes(b"modified")
        with self.assertRaises(ValueError):
            validate_receipt(request, receipt)

    def test_empty_week_can_be_finalized(self):
        run_dir = self.collect([])
        self.finalize(run_dir)
        self.assertTrue((run_dir / "Policy_Weekly.xlsx").exists())

    def test_scope_filtered_metadata_is_saved_without_ai_or_body_failure(self):
        item = {"title": "项目申报通知", "detail_url": "https://example.com/unconfirmed.html",
                "decision": "Needs review", "reason": "ENERGY_TOPIC_UNCONFIRMED", "body_fetched": False}
        statuses = [{"agency": "TEST", "status": "ScopeReview", "scope_filtered": [item]}]
        with patch("policy_agent.cli.WeeklyCollector.collect", return_value=([], statuses)):
            code = main(["--installation", str(self.installation), "--config", str(self.sources), "--out-dir", str(self.out)])
        self.assertEqual(code, 0)
        run_dir = next(self.out.glob("*/run.json")).parent
        self.assertEqual(read_json(run_dir / "scope_filtered.json")["records"], [item])
        self.assertEqual(read_json(run_dir / "analysis_request.json")["records"], [])
        self.assertEqual(read_json(run_dir / "failed_details.json")["records"], [])
        self.assertEqual(read_json(run_dir / "run.json")["meta"]["scope_review_records"], 1)


if __name__ == "__main__":
    unittest.main()
