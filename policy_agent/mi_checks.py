"""Deterministic MI reference document and saved-workbook consistency checks."""
from pathlib import Path
import hashlib
import json
from zipfile import ZipFile
from urllib.parse import urlsplit
from .paths import PACKAGE_ROOT, local_path
from .taxonomy import dictionary, agency_by_code, split_labels, validate_classification

REFERENCE_DOCUMENT = PACKAGE_ROOT / "docs" / "MI_DICTIONARY.md"


def render_dictionary(data):
    def text(value):
        return str(value).replace("|", "\\|").replace("\n", " ")

    def table(headers, rows):
        return ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |",
                *["| " + " | ".join(text(v) for v in row) + " |" for row in rows], ""]

    digest = hashlib.sha256(json.dumps(data, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")).hexdigest()
    lines = ["---", "name: local-policy-dictionary", "description: 本地机构译名、分类标签和输出字段的核对基准。", "---", "",
             "# 本地机构与政策标签字典", "",
             "本文从本地 JSON 生成，不应上传到公开仓库。检查时必须逐条核对与用户参考表是否一致。",
             f"参考文件：`{data.get('reference_file', '')}`。",
             f"参考文件 SHA256：`{data.get('reference_sha256', '')}`。",
             f"规范化字典 SHA256：`{digest}`。", "",
             "核对机构代码、中英文译名、机关与下属办公厅、联合发文顺序，以及每个行业和标签的配对。",
             "参考表有嵌入图片时须人工逐条核对译名，并登记图片部件和 SHA256；单元格检查不能代替图片核对。",
             "生成表必须与字典逐项一致；Full_Text_Link 和 Source_List_URL 必须具备真实 Excel 超链接属性。",
             "未知机构不得猜译，正文缺失不得填作政策摘要。不一致时停止定稿和邮件交付。",
             "未提供参考原表时仅核验本地资产和生成表，不能声称重新核验了原表。", "",
             "## Issuing Agency", ""]
    lines += table(["Code", "English name", "中文名称"],
                   [[a["code"], a["name"], a["chinese_name"]] for a in data["agencies"]])
    lines += ["## 下属单位与历史机构", ""]
    lines += table(["Code", "English name", "中文名称"],
                   [[a["code"], a["name"], a["chinese_name"]] for a in data["issuer_units"]])
    lines += ["## 已人工核对的联合发文示例", ""]
    lines += table(["序号", "中文", "英文"],
                   [[a["number"], a["chinese_name"], a["english_name"]] for a in data["issuing_agency_examples"]])
    lines += ["## Supported_Industry / Policy Tags", ""]
    lines += table(["Category", "Tag", "参考位置"],
                   [[p["category"], p["tag"], p["reference_cell"]] for p in data["policy_tags"]])
    lines += ["## 输出字段", ""]
    lines += table(["序号", "字段", "规则"],
                   [[i, f["name"], f["rule"]] for i, f in enumerate(data["export_fields"], 1)])
    lines += ["## 参考字段别名", ""]
    lines += table(["参考字段", "输出字段"], list(data["reference_field_aliases"].items()))
    lines += ["运行 scripts/check_mi_dictionary.py；有原表时加 --reference，有生成表时加 --workbook。", ""]
    return "\n".join(lines)


def check_dictionary_assets(root=None, verify_manifest=False):
    default_root = root is None
    root = Path(root) if root else PACKAGE_ROOT
    data = json.loads((root / "policy_agent/data/mi_dictionary.json").read_text(encoding="utf-8"))
    if default_root and dictionary() != data:
        raise ValueError("运行期间 MI 字典发生变化，请先核对原表并重启当前任务")
    expected = render_dictionary(data)
    if (root / "docs/MI_DICTIONARY.md").read_text(encoding="utf-8") != expected:
        raise ValueError("MI_DICTIONARY.md 与程序字典不一致，请重新生成并与 MI 表核对")
    agencies = data["agencies"] + data["issuer_units"]
    codes = [a["code"] for a in agencies]
    if len(codes) != len(set(codes)) or any(not a.get(k) for a in agencies for k in ("code", "name", "chinese_name")):
        raise ValueError("MI 机构字典存在重复代码或空值")
    pairs = [(p["category"], p["tag"]) for p in data["policy_tags"]]
    if not pairs or len(pairs) != len(set(pairs)):
        raise ValueError("MI 分类配对为空或重复")
    from .export import HEADERS
    if [f["name"] for f in data["export_fields"]] != HEADERS:
        raise ValueError("本地字典的导出字段和程序固定格式不一致")
    by_code = {a["code"]: a for a in agencies}
    for example in data["issuing_agency_examples"]:
        if ", ".join(by_code[code]["name"] for code in example["codes"]) != example["english_name"]:
            raise ValueError("Issuing Agency 联合发文示例与机构译名字典不一致")
    manifest_checked = False
    if verify_manifest:
        manifest = json.loads((root / "PACKAGE_MANIFEST.json").read_text(encoding="utf-8"))
        entries = {item["path"]: item for item in manifest["files"]}
        for relative in ("policy_agent/data/mi_dictionary.json", "docs/MI_DICTIONARY.md"):
            if relative not in entries or hashlib.sha256((root / relative).read_bytes()).hexdigest() != entries[relative]["sha256"]:
                raise ValueError(f"迁移字典缺失或与发布清单不一致：{relative}")
        manifest_checked = True
    return {"assets_consistent": True, "agency_records": len(data["agencies"]),
            "issuer_units": len(data["issuer_units"]), "tag_pairs": len(pairs), "manifest_checked": manifest_checked}


def check_reference(path):
    import openpyxl
    data = dictionary()
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        agencies = [{"code": r[0], "name": r[1], "chinese_name": r[2]}
                    for r in list(wb["Agency abb."].values)[1:] if r[0]]
        expected_agencies = [{k: a[k] for k in ("code", "name", "chinese_name")} for a in data["agencies"]]
        if agencies != expected_agencies:
            raise ValueError("原 MI Agency abb. 的代码、中英文名称或顺序与独立字典不一致")
        ws = wb["Policy Tags"]
        tags = []
        for col in range(2, ws.max_column + 1):
            category = ws.cell(1, col).value
            for row in range(2, ws.max_row + 1):
                tag = ws.cell(row, col).value
                if tag:
                    if not isinstance(tag, str) or not isinstance(category, str):
                        raise ValueError("原 MI 分类标签含无效列头或非文本标签")
                    tags.append({"category": category.strip(), "tag": tag.strip(), "reference_cell": ws.cell(row, col).coordinate})
        if tags != data["policy_tags"]:
            raise ValueError("原 MI Policy Tags 的类别、标签、配对或单元格位置与独立字典不一致")
    finally:
        wb.close()
    with ZipFile(path) as archive:
        for entry in data["issuing_agency_images"]:
            if entry["part"] not in archive.namelist() or hashlib.sha256(archive.read(entry["part"])).hexdigest() != entry["sha256"]:
                raise ValueError("Issuing Agency 图片变化或丢失，必须人工核对译名后更新字典")
    return {"reference_checked": True, "reference_dictionary_consistent": True,
            "reference_file_hash_matches": hashlib.sha256(Path(path).read_bytes()).hexdigest() == data["reference_sha256"],
            "embedded_images_unchanged": True}


def check_workbook(path):
    import openpyxl
    data = dictionary()
    wb = openpyxl.load_workbook(local_path(path, "核对工作簿", must_exist=True), read_only=False)
    try:
        ws = wb["Weekly_Policies"]
        headers = [c.value for c in ws[1]]
        expected = [f["name"] for f in data["export_fields"]]
        if headers != expected:
            raise ValueError("生成表的 Weekly_Policies 字段及顺序与独立字典不一致")
        columns = {name: n + 1 for n, name in enumerate(headers)}
        for row in range(2, ws.max_row + 1):
            def value(name):
                return ws.cell(row, columns[name]).value or ""
            codes = split_labels(value("Issuing_Agency"))
            entries = [agency_by_code(code) for code in codes]
            if not entries or " / ".join(a["code"] for a in entries) != value("Issuing_Agency"):
                raise ValueError(f"生成表第{row}行机构代码不规范")
            if ", ".join(a["name"] for a in entries) != value("Issuing_Agency_Name"):
                raise ValueError(f"生成表第{row}行英文机构译名与字典不一致")
            industry, tags = validate_classification(value("Supported_Industry"), value("Policy_Tags"))
            if (industry, tags) != (value("Supported_Industry"), value("Policy_Tags")):
                raise ValueError(f"生成表第{row}行分类和标签格式与字典不一致")
            for field in ("Full_Text_Link", "Source_List_URL"):
                cell = ws.cell(row, columns[field])
                if not cell.value and field == "Source_List_URL":
                    continue
                url = cell.value
                if not isinstance(url, str) or urlsplit(url).scheme not in {"http", "https"} or not urlsplit(url).netloc:
                    raise ValueError(f"生成表第{row}行{field} URL 无效")
                if not cell.hyperlink or cell.hyperlink.target != url:
                    raise ValueError(f"生成表第{row}行{field}是普通文本或超链接目标不一致")
        if list(wb["Tag_Dictionary"].values) != [("Category", "Tag"), *[(p["category"], p["tag"]) for p in data["policy_tags"]]]:
            raise ValueError("生成表 Tag_Dictionary 与独立字典不一致")
        return {"workbook_checked": True, "policies_checked": max(ws.max_row - 1, 0), "hyperlinks_checked": True}
    finally:
        wb.close()
