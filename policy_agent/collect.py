"""Fetch verified policy bodies; reachability and column discovery are not success."""
from io import BytesIO
from time import perf_counter
from urllib.parse import urlsplit, urlunsplit
from zipfile import ZipFile
from xml.etree import ElementTree
from threading import RLock
from posixpath import normpath
from base64 import b64decode
import hashlib
import json

from .http import fetch, FetchResult
from .models import Source
from .parse import (select_policy_links, extract_detail, extract_text_detail, normalize_url,
                    discover_entries, attachment_links, body_image_links, requires_policy_attachment, EXCLUDE_WORDS)
from .paths import local_path
from .runtime import configure
from .scope import energy_policy_candidate


def browser_fetch(url):
    configure()
    started = perf_counter()
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as manager:
            browser = manager.chromium.launch(headless=True)
            page = browser.new_page()
            navigation_responses = []
            def record_response(response):
                if response.request.is_navigation_request() and response.request.frame == page.main_frame:
                    navigation_responses.append((response.url, response.status))
            page.on("response", record_response)
            response = page.goto(url, wait_until="domcontentloaded", timeout=30000)
            try:
                page.wait_for_load_state("networkidle", timeout=8000)
            except Exception:
                pass
            page.wait_for_timeout(1500 if response and response.status >= 400 else 500)
            html = page.content().encode("utf-8")
            final_status = navigation_responses[-1][1] if navigation_responses else response.status if response else None
            result = FetchResult(url, page.url, final_status,
                                 "text/html", html.decode("utf-8"), html,
                                 int((perf_counter() - started) * 1000))
            browser.close()
            return result
    except Exception as exc:
        return FetchResult(url, url, None, "", "", b"", int((perf_counter() - started) * 1000), str(exc))


def browser_binary_fetch(url, referer):
    """Read an allowed public attachment in the official page's browser context."""
    configure()
    started = perf_counter()
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as manager:
            browser = manager.chromium.launch(headless=True)
            page = browser.new_page()
            page.goto(referer, wait_until="domcontentloaded", timeout=30000)
            try:
                page.wait_for_load_state('networkidle', timeout=8000)
            except Exception:
                pass
            page.wait_for_timeout(300)
            payload = page.evaluate("""async url => {
                const response = await fetch(url, {signal: AbortSignal.timeout(30000)});
                const buffer = new Uint8Array(await response.arrayBuffer());
                if (buffer.length > 100000000) throw new Error('ASSET_SIZE_LIMIT_EXCEEDED');
                let binary = '';
                for (let i=0; i<buffer.length; i+=32768)
                    binary += String.fromCharCode(...buffer.subarray(i, i+32768));
                return {url: response.url, status: response.status,
                    type: response.headers.get('content-type') || '', data: btoa(binary)};
            }""", url)
            browser.close()
        raw = b64decode(payload["data"], validate=True)
        status = payload["status"]
        return FetchResult(url, payload["url"], status, payload["type"], "", raw,
                           int((perf_counter() - started) * 1000), "" if 200 <= status < 300 else f"HTTP_{status}")
    except Exception as exc:
        return FetchResult(url, url, None, "", "", b"", int((perf_counter() - started) * 1000), str(exc))


_document_lock = RLock()


def document_text(response, processing_dir=None):
    # PDFium shares native process state; nested archive parsing uses the same lock.
    with _document_lock:
        return _document_text(response, processing_dir)


def _document_text(response, processing_dir=None):
    suffix = urlsplit(response.final_url).path.lower()
    content_type = response.content_type.lower()
    if response.content.lstrip().startswith(b"%PDF-") or "pdf" in content_type or suffix.endswith(".pdf"):
        import pypdfium2
        from .ocr import recognize
        pages, evidence = [], []
        with pypdfium2.PdfDocument(response.content) as pdf:
            if len(pdf) > 1000:
                raise ValueError("PDF_PAGE_LIMIT_EXCEEDED")
            for index in range(len(pdf)):
                page = pdf[index]
                textpage = page.get_textpage()
                text = textpage.get_text_bounded()
                textpage.close()
                images = list(page.get_objects(filter=[3]))
                area = page.get_width() * page.get_height()
                large_image = any((bounds[2] - bounds[0]) * (bounds[3] - bounds[1]) > area * .50
                                  for bounds in (item.get_bounds() for item in images))
                if images and (len(text.strip()) < 80 or (large_image and len(text.strip()) < 200)):
                    if processing_dir is None:
                        raise ValueError("PDF_OCR_REQUIRED")
                    processing_dir.mkdir(parents=True, exist_ok=True)
                    image_path = processing_dir / f"page_{index + 1:04d}.png"
                    bitmap = page.render(scale=3)
                    bitmap.to_pil().save(image_path)
                    bitmap.close()
                    result = recognize(image_path, processing_dir / f"page_{index + 1:04d}.ocr.json", min_characters=1)
                    if not result["valid"]:
                        raise ValueError(f"PDF_PAGE_OCR_FAILED: page {index + 1}: " + result["error"])
                    text = result["text"]
                    evidence.append({"page": index + 1, "method": "ocr", "image_path": str(image_path), "mean_confidence": result["mean_confidence"]})
                else:
                    evidence.append({"page": index + 1, "method": "pdf_text"})
                pages.append(text)
                page.close()
        if processing_dir is not None:
            processing_dir.mkdir(parents=True, exist_ok=True)
            (processing_dir / "pages.json").write_text(json.dumps({"page_count": len(pages), "pages": [{**item, "text": text} for item, text in zip(evidence, pages)]}, ensure_ascii=False, indent=2), encoding="utf-8")
        return "\n".join(pages), ".pdf"
    if response.content.startswith(b"PK"):
        with ZipFile(BytesIO(response.content)) as archive:
            if "word/document.xml" in archive.namelist():
                root = ElementTree.fromstring(archive.read("word/document.xml"))
                text = "\n".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
                references = [value for node in root.iter() for name, value in node.attrib.items() if name.endswith("}embed") or (node.tag.endswith("}imagedata") and name.endswith("}id"))]
                if any(node.tag.endswith("}object") for node in root.iter()):
                    raise ValueError("DOCX_EMBEDDED_OBJECT_REQUIRES_EXTRACTION")
                if references:
                    from .ocr import recognize
                    if processing_dir is None:
                        raise ValueError("DOCX_IMAGE_OCR_REQUIRED")
                    processing_dir.mkdir(parents=True, exist_ok=True)
                    relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
                    targets = {node.get("Id"): node.get("Target") for node in relationships if node.get("TargetMode") != "External"}
                    for index, reference in enumerate(dict.fromkeys(references)):
                        target = normpath("word/" + targets.get(reference, ""))
                        if not target.startswith("word/media/") or target not in archive.namelist():
                            raise ValueError("DOCX_BODY_IMAGE_MISSING")
                        image_path = processing_dir / f"body_image_{index + 1:03d}.img"
                        image_path.write_bytes(archive.read(target))
                        result = recognize(image_path, processing_dir / f"body_image_{index + 1:03d}.ocr.json", min_characters=1)
                        if not result["valid"]:
                            raise ValueError("DOCX_IMAGE_OCR_INCOMPLETE: " + result["error"])
                        text += "\n【正文图片】\n" + result["text"]
                return text, ".docx"
            if "xl/workbook.xml" in archive.namelist():
                from openpyxl import load_workbook
                workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
                formulas = load_workbook(BytesIO(response.content), read_only=True, data_only=False)
                try:
                    if sum(sheet.max_row * sheet.max_column for sheet in workbook.worksheets) > 500_000:
                        raise ValueError("SPREADSHEET_SIZE_LIMIT_EXCEEDED")
                    for cached, original in zip(workbook.worksheets, formulas.worksheets):
                        for cached_row, original_row in zip(cached.iter_rows(), original.iter_rows()):
                            if any(cell.data_type == "f" and value.value is None for cell, value in zip(original_row, cached_row)):
                                raise ValueError("SPREADSHEET_FORMULA_VALUE_MISSING")
                    if any(name.startswith("xl/media/") for name in archive.namelist()):
                        raise ValueError("SPREADSHEET_BODY_IMAGE_REQUIRES_EXTRACTION")
                except Exception:
                    workbook.close()
                    raise
                finally:
                    formulas.close()
                text = "\n".join("【" + sheet.title + "】\n" + "\n".join("\t".join(str(cell) if cell is not None else "" for cell in row) for row in sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets)
                workbook.close()
                return text, ".xlsx"
            members = [member for member in archive.infolist() if not member.is_dir() and member.filename.lower().endswith((".doc", ".wps", ".docx", ".pdf", ".xlsx", ".xls"))]
            if members and len(members) <= 30 and sum(member.file_size for member in members) <= 100_000_000:
                texts = []
                for index, member in enumerate(members):
                    raw = archive.read(member)
                    child_dir = processing_dir / f"member_{index + 1:03d}" if processing_dir is not None else None
                    if child_dir is not None:
                        child_dir.mkdir(parents=True, exist_ok=True)
                        (child_dir / ("original" + member.filename[member.filename.rfind("."):])).write_bytes(raw)
                    child = FetchResult(response.url, response.final_url.split("?")[0] + "/" + member.filename, 200, "", "", raw, 0)
                    text, _ = document_text(child, child_dir)
                    if not text or len(text.strip()) < 40:
                        raise ValueError("ZIP_MEMBER_TEXT_INCOMPLETE: " + member.filename)
                    texts.append("【" + member.filename + "】\n" + text)
                return "\n\n".join(texts), ".zip"
        raise ValueError("UNSUPPORTED_DOCUMENT_CONTAINER")
    if response.content.startswith(b"\xd0\xcf\x11\xe0") or suffix.endswith((".doc", ".wps")):
        import olefile
        with olefile.OleFileIO(BytesIO(response.content)) as container:
            spreadsheet = container.exists("Workbook") or container.exists("Book")
        if spreadsheet:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=response.content)
            if sum(sheet.nrows * sheet.ncols for sheet in workbook.sheets()) > 500_000:
                raise ValueError("SPREADSHEET_SIZE_LIMIT_EXCEEDED")
            return "\n".join("【" + sheet.name + "】\n" + "\n".join("\t".join(str(value) for value in sheet.row_values(index)) for index in range(sheet.nrows)) for sheet in workbook.sheets()), ".xls"
        from .legacy_doc import extract_legacy_document
        text, metadata = extract_legacy_document(response.content)
        if processing_dir is not None:
            processing_dir.mkdir(parents=True, exist_ok=True)
            (processing_dir / "legacy_text.json").write_text(json.dumps({**metadata, "text": text}, ensure_ascii=False, indent=2), encoding="utf-8")
        return text, ".doc"
    if suffix.endswith(".docx") or "wordprocessingml" in content_type:
        raise ValueError("INVALID_DOCX_RESPONSE")
    return None, ".html"


class WeeklyCollector:
    def __init__(self, raw_dir, timeout=20, retries=2, use_browser=False):
        self.raw_dir = local_path(raw_dir, "证据目录")
        self.raw_dir.mkdir(parents=True, exist_ok=True)
        self.timeout, self.retries, self.use_browser = timeout, retries, use_browser
        self.discovery_limit = 6
        self.identity_cache = {}

    def _evidence(self, source, url, content, suffix=".html", phase="detail"):
        key = hashlib.sha256((phase + "|" + url).encode("utf-8")).hexdigest()[:16]
        folder = local_path(self.raw_dir / source.code.replace("/", "_").replace(" ", "_"), "证据目录")
        folder.mkdir(parents=True, exist_ok=True)
        path = local_path(folder / (key + suffix), "证据文件")
        path.write_bytes(content)
        return str(path)

    def _links(self, source, page, limit):
        return select_policy_links(page.text, page.final_url, limit,
                                   detail_pattern=source.policy_detail_pattern,
                                   allowed_hosts=source.allowed_hosts)

    def _asset(self, source, url, referer, attempted):
        asset = fetch(url, self.timeout, self.retries)
        if self.use_browser and (not asset.status or asset.status >= 400 or "text/html" in asset.content_type.lower()):
            attempted.append({"method": "http_asset", "url": asset.final_url, "http_status": asset.status,
                              "error": asset.error or "ASSET_HTML_RESPONSE",
                              "evidence_path": self._evidence(source, asset.final_url, asset.content, ".bin", "asset_http") if asset.content else ""})
            rendered = browser_binary_fetch(url, referer)
            attempted.append({"method": "browser_asset", "url": rendered.final_url, "http_status": rendered.status,
                              "error": rendered.error,
                              "evidence_path": self._evidence(source, rendered.final_url, rendered.content, ".bin", "asset_browser") if rendered.content else ""})
            if rendered.status and 200 <= rendered.status < 300 and rendered.content:
                asset = rendered
        return asset

    def _article_page(self, source, page, attempted):
        if "#news-details" not in page.text:
            return page
        from .dynamic import hydrate_literal_article
        hydrated = hydrate_literal_article(page)
        if hydrated.text != page.text:
            attempted.append({"method": "literal_article", "url": page.final_url, "http_status": page.status,
                              "evidence_path": self._evidence(source, page.final_url, page.content, ".html", "literal_article_original")})
        return hydrated

    def _parse(self, source, page, title):
        key = hashlib.sha256(page.final_url.encode("utf-8")).hexdigest()[:16]
        processing = local_path(self.raw_dir / source.code.replace("/", "_").replace(" ", "_") / "processing" / key, "正文解析证据")
        try:
            text, suffix = document_text(page, processing)
            parsed = extract_text_detail(text, page.final_url, title) if text is not None else extract_detail(
                page.text, page.final_url, title, source.body_selectors)
            parsed["processing_evidence"] = [str(path) for path in processing.rglob("*") if path.is_file()] if processing.exists() else []
            return parsed, suffix
        except Exception as exc:
            parsed = extract_text_detail("", page.final_url, title)
            parsed["validation_error"] = "DOCUMENT_PARSE_FAILED: " + str(exc)[:180]
            return parsed, ".pdf" if "pdf" in page.content_type.lower() else ".bin"

    def _detail(self, source, link, list_url):
        attempted = []
        path = urlsplit(link["url"]).path.lower()
        direct_document = path.endswith((".pdf", ".doc", ".wps", ".docx", ".xls", ".xlsx", ".zip")) or "/gwk/resource/file?" in link["url"]
        page = self._asset(source, link["url"], list_url, attempted) if direct_document else fetch(link["url"], self.timeout, self.retries)
        if not direct_document and link["url"].startswith("http://") and (not page.status or page.status >= 400):
            from .navigation import secure_alternative
            attempted.append({"method": "http_detail", "url": page.final_url, "http_status": page.status,
                              "error": page.error,
                              "evidence_path": self._evidence(source, page.final_url, page.content, ".html", "detail_http") if page.content else ""})
            secure = fetch(secure_alternative(link["url"]), self.timeout, self.retries)
            if secure.status and 200 <= secure.status < 300:
                page = secure
            else:
                attempted.append({"method": "https_detail", "url": secure.final_url, "http_status": secure.status,
                                  "error": secure.error,
                                  "evidence_path": self._evidence(source, secure.final_url, secure.content, ".html", "detail_https") if secure.content else ""})
        page = self._article_page(source, page, attempted)
        parsed, suffix = self._parse(source, page, link["text"])
        evidence = self._evidence(source, page.final_url, page.content, suffix) if page.content else ""
        valid_http = lambda response: bool(response.status and 200 <= response.status < 300)
        attempted.append({"method": "http", "url": page.final_url, "http_status": page.status,
                          "error": page.error or parsed.get("validation_error"), "evidence_path": evidence})
        method = "http"
        context_html = False
        # Both HTTP errors and 200 responses without an article need a rendered retry.
        host = urlsplit(source.website).hostname or ""
        attachments = attachment_links(page.text, page.final_url, host, source.allowed_hosts) if suffix == ".html" else []
        if self.use_browser and suffix == ".html" and (not valid_http(page) or not parsed["body_valid"] or
                (requires_policy_attachment(parsed) and not attachments)):
            rendered = self._article_page(source, browser_fetch(link["url"]), attempted)
            rendered_parsed, rendered_suffix = self._parse(source, rendered, link["text"])
            rendered_evidence = self._evidence(source, rendered.final_url, rendered.content, rendered_suffix, "rendered") if rendered.content else ""
            attempted.append({"method": "browser", "url": rendered.final_url, "http_status": rendered.status,
                              "error": rendered.error or rendered_parsed.get("validation_error"), "evidence_path": rendered_evidence})
            if valid_http(rendered) and (rendered_parsed["body_valid"] or not parsed["body_valid"]):
                page, parsed, evidence, method = rendered, rendered_parsed, rendered_evidence, "browser"
                attachments = attachment_links(page.text, page.final_url, host, source.allowed_hosts)
        if self.use_browser and suffix == '.html' and not valid_http(page):
            from .navigation import secure_alternative
            context = browser_binary_fetch(secure_alternative(link['url']) if list_url.startswith('https:') else link['url'], list_url)
            if context.content and 'html' in context.content_type.lower():
                from bs4 import UnicodeDammit
                context.text = UnicodeDammit(context.content, is_html=True).unicode_markup or ''
            context = self._article_page(source, context, attempted)
            context_parsed, context_suffix = self._parse(source, context, link['text'])
            context_evidence = self._evidence(source, context.final_url, context.content, context_suffix, 'browser_context') if context.content else ''
            attempted.append({'method':'browser_context','url':context.final_url,'http_status':context.status,
                              'error':context.error or context_parsed.get('validation_error'),'evidence_path':context_evidence})
            if valid_http(context) and context_suffix == '.html':
                page, parsed, evidence, method = context, context_parsed, context_evidence, 'browser_context'
                context_html = True
                attachments = attachment_links(page.text, page.final_url, host, source.allowed_hosts)
        attachment_paths = []
        processing_paths = list(parsed.get("processing_evidence", []))
        image_urls = body_image_links(page.text, page.final_url, host, source.allowed_hosts, source.body_selectors) if suffix == ".html" else []
        image_texts, image_errors = [], []
        for index, image_url in enumerate(image_urls[:40]):
            image = self._asset(source, image_url, page.final_url, attempted)
            image_path = self._evidence(source, image_url, image.content, ".img", "body_image") if image.content else ""
            image_error = image.error
            ocr_path = ""
            if valid_http(image) and image_path:
                try:
                    from .ocr import recognize
                    ocr_path = str(local_path(image_path + ".ocr.json", "OCR 证据"))
                    result = recognize(image_path, ocr_path, min_characters=1)
                    if result["valid"]:
                        image_texts.append(f"【正文图片 {index + 1}】 " + result["text"])
                        processing_paths.extend([image_path, ocr_path])
                    else:
                        image_error = result["error"]
                except Exception as exc:
                    image_error = "OCR_FAILED: " + str(exc)[:180]
            else:
                image_error = image_error or "IMAGE_DOWNLOAD_FAILED"
            attempted.append({"method": "image_ocr", "url": image_url, "http_status": image.status, "error": image_error, "evidence_path": image_path, "ocr_path": ocr_path})
            if image_error:
                image_errors.append(image_error)
        if len(image_urls) > 40:
            image_errors.append("BODY_IMAGE_LIMIT_EXCEEDED")
        if image_texts and not image_errors:
            parsed["body_text"] = "\n\n".join([parsed["body_text"]] + image_texts).strip()
            parsed["body_valid"] = len(parsed["body_text"]) >= 80
            parsed["validation_error"] = "" if parsed["body_valid"] else "IMAGE_BODY_TEXT_INCOMPLETE"
            parsed["body_selector"] += " + image OCR"
            method = "image_ocr"
        elif image_errors:
            parsed["body_valid"], parsed["validation_error"] = False, "BODY_IMAGE_OCR_INCOMPLETE: " + "; ".join(image_errors)
        required = suffix == ".html" and ((not parsed["body_valid"] and bool(attachments)) or requires_policy_attachment(parsed))
        body_sources = [{"url": page.final_url, "evidence_path": evidence, "kind": "html" if suffix == ".html" else "document"}]
        body_sources.extend({"url": item["url"], "evidence_path": item["evidence_path"], "kind": "html_original"}
                            for item in attempted if item.get("method") == "literal_article" and item.get("evidence_path"))
        attachment_bodies, attachment_errors = [], []
        if valid_http(page) and required:
            for attachment in attachments[:10]:
                document = self._asset(source, attachment["url"], page.final_url, attempted)
                doc_parsed, doc_suffix = self._parse(source, document, link["text"])
                doc_path = self._evidence(source, document.final_url, document.content, doc_suffix, "attachment") if document.content else ""
                attachment_paths.append(doc_path)
                attempted.append({"method": "attachment", "url": document.final_url, "http_status": document.status,
                                  "error": document.error or doc_parsed.get("validation_error"), "evidence_path": doc_path})
                if valid_http(document) and doc_parsed["body_valid"]:
                    processing_paths.extend(doc_parsed.get("processing_evidence", []))
                    attachment_bodies.append("【附件：" + attachment["text"] + "】 " + doc_parsed["body_text"])
                    body_sources.append({"url": document.final_url, "evidence_path": doc_path, "kind": "document"})
                else:
                    attachment_errors.append(document.error or doc_parsed.get("validation_error") or "ATTACHMENT_DOWNLOAD_FAILED")
            if not attachments:
                attachment_errors.append("POLICY_ATTACHMENT_LINK_MISSING")
            elif len(attachments) > 10:
                attachment_errors.append("ATTACHMENT_LIMIT_EXCEEDED")
            if attachment_bodies and not attachment_errors:
                parsed["body_text"] = "\n\n".join([parsed["body_text"]] + attachment_bodies).strip()
                parsed["abstract"] = parsed["body_text"][:1000]
                parsed["body_valid"] = True
                parsed["body_selector"] += " + policy attachments"
                method = "attachment"
            else:
                parsed["body_valid"] = False
                parsed["validation_error"] = "POLICY_ATTACHMENT_INCOMPLETE: " + "; ".join(attachment_errors)
        final_url = normalize_url(page.final_url)
        if context_html:
            requested, returned = urlsplit(link['url']), urlsplit(final_url)
            if requested.hostname == returned.hostname and requested.path == returned.path:
                final_url = urlunsplit((returned.scheme, returned.netloc, returned.path, requested.query, ''))
        from .navigation import same_page
        column_returned = same_page(final_url, list_url) or same_page(final_url, source.website)
        valid = valid_http(page) and parsed["body_valid"] and bool(evidence) and not column_returned
        error = "" if valid else ("COLUMN_PAGE" if column_returned else
                                 page.error or parsed.get("validation_error") or "HTTP_DETAIL_FAILED")
        if valid and EXCLUDE_WORDS.search(parsed["title"]):
            valid, error = False, "NOT_POLICY_DOCUMENT"
        return {"agency": source.code, "agency_name": source.chinese_name or source.name,
                "list_url": list_url, "detail_url": final_url, "title": parsed["title"],
                "response_url": page.final_url,
                "issue_date": (link.get("date") or parsed["issue_date"]) if suffix != ".html" else (parsed["issue_date"] or link.get("date", "")), "policy_number": parsed["policy_number"],
                "body_text": parsed["body_text"] if valid else "", "abstract": parsed["abstract"] if valid else "",
                "http_status": page.status or "", "collector_method": method,
                "review_status": "Collected" if valid else "Failed", "body_valid": valid,
                "body_selector": parsed.get("body_selector", ""), "evidence_path": evidence,
                "body_sha256": hashlib.sha256(parsed["body_text"].encode("utf-8")).hexdigest() if valid else "",
                "body_sources": body_sources, "processing_evidence": processing_paths,
                "attachment_evidence": attachment_paths, "attempts": attempted, "error": error}

    def collect_source(self, source, max_items=None, since=None, until=None):
        from datetime import date
        from .navigation import pagination_links, secure_alternative, authority_matches, same_page, energy_entries
        limit = max_items or source.max_items
        if not source.enabled:
            return [], {"agency": source.code, "status": "Disabled", "list_url": source.list_url, "candidates": 0, "error": "disabled"}
        roots = list(dict.fromkeys(source.entry_urls + [source.list_url or source.website]))
        queue = [(url, 0, "entry") for url in roots]
        visited, candidates, discovery = set(), {}, []
        scope_filtered = {}
        list_origins = {}
        pagination_count, boundary_seen, exhausted, limited = 0, False, False, False
        while queue and len(visited) < self.discovery_limit:
            url, depth, kind = queue.pop(0)
            key = normalize_url(url)
            if key in visited:
                continue
            visited.add(key)
            page = fetch(url, self.timeout, self.retries)
            if (not page.status or page.status >= 400) and url.startswith("http://"):
                secure = fetch(secure_alternative(url), self.timeout, self.retries)
                if secure.status and 200 <= secure.status < 300:
                    page = secure
            from .dynamic import hydrate_datastore, literal_list_redirect
            list_original_evidence = ''
            hydrated = hydrate_datastore(page, source.allowed_hosts) if page.status == 200 else page
            if hydrated is not page:
                list_original_evidence = self._evidence(source, page.final_url, page.content, phase='list_original')
                page = hydrated
            redirect_target = literal_list_redirect(page, source.allowed_hosts) if page.status == 200 else ''
            if redirect_target and normalize_url(redirect_target) not in visited:
                queue.insert(0, (redirect_target, depth, 'redirect'))
            links = self._links(source, page, max(limit + 1, 100)) if page.status and 200 <= page.status < 300 else []
            method = "http"
            page_evidence = self._evidence(source, page.final_url, page.content, phase="list") if page.content else ""
            api_attempts = []
            shanghai_entry = urlsplit(page.final_url).hostname == "www.shanghai.gov.cn" and urlsplit(page.final_url).path.startswith("/zhengce/")
            if page.status == 200 and ("front/page/build/unit" in page.text or shanghai_entry):
                try:
                    from .dynamic import unit_pages, shanghai_policy_pages
                    adapter = shanghai_policy_pages if shanghai_entry else unit_pages
                    hydrated, responses, api_complete = adapter(page, self.timeout, source.pagination_pages if since else 1, since)
                    for api_index, api_response in enumerate(responses):
                        api_attempts.append({"url": api_response.final_url, "http_status": api_response.status,
                                             "evidence_path": self._evidence(source, api_response.final_url, api_response.content, ".json", f"unit_api_{api_index}")})
                    page, method = hydrated, "unit_api"
                    links = self._links(source, page, max(limit + 1, 100))
                    if since and not api_complete:
                        limited = True
                    page_evidence = self._evidence(source, page.final_url, page.content, phase="hydrated_list")
                except Exception as exc:
                    api_attempts.append({"error": "UNIT_API_FAILED: " + str(exc)[:200]})
            if not links and self.use_browser and not redirect_target:
                rendered = browser_fetch(page.final_url if page.status and page.status < 400 else url)
                if rendered.status and 200 <= rendered.status < 300:
                    rendered_links = self._links(source, rendered, max(limit + 1, 100))
                    if rendered_links or not page.status or page.status >= 400:
                        page, links, method = rendered, rendered_links, "browser"
                        page_evidence = self._evidence(source, page.final_url, page.content, phase="rendered_list")
            if not links and self.use_browser and (not page.status or page.status >= 400):
                context = browser_binary_fetch(url, source.website)
                context_evidence = self._evidence(source, context.final_url, context.content, phase='browser_context_list') if context.content else ''
                api_attempts.append({'url':context.final_url,'http_status':context.status,'error':context.error,
                                     'method':'browser_context_list','evidence_path':context_evidence})
                if context.status == 200 and 'html' in context.content_type.lower():
                    from bs4 import UnicodeDammit
                    context.text = UnicodeDammit(context.content, is_html=True).unicode_markup or ''
                    requested, returned = urlsplit(url), urlsplit(context.final_url)
                    if requested.hostname == returned.hostname and requested.path == returned.path:
                        context.final_url = urlunsplit((returned.scheme, returned.netloc, returned.path, requested.query, ''))
                    context = hydrate_datastore(context, source.allowed_hosts)
                    page, links, method = context, self._links(source, context, max(limit + 1, 100)), 'browser_context'
                    page_evidence = self._evidence(source, page.final_url, page.content, phase='browser_context_list_parsed')
            homepage = same_page(page.final_url, source.website)
            agency_ok = authority_matches(source, page.text, page.final_url)
            authority_evidence = ""
            energy_portal_inherited = False
            if (not agency_ok and source.code.split()[-1] in {"EA", "EB"}
                    and urlsplit(page.final_url).hostname == urlsplit(source.website).hostname):
                from dataclasses import replace
                drc = replace(source, code=source.code.split()[0] + " DRC")
                if authority_matches(drc, page.text, page.final_url):
                    if source.code not in self.identity_cache:
                        home = fetch(source.website, self.timeout, self.retries)
                        if (not home.status or home.status >= 400) and source.website.startswith("http://"):
                            home = fetch(secure_alternative(source.website), self.timeout, self.retries)
                        self.identity_cache[source.code] = home
                    home = self.identity_cache[source.code]
                    if (home.status == 200 and urlsplit(home.final_url).hostname == urlsplit(source.website).hostname
                            and authority_matches(source, home.text, home.final_url)):
                        agency_ok = energy_portal_inherited = True
                        authority_evidence = self._evidence(source, home.final_url, home.content, phase="energy_entry_identity")
            if not agency_ok and source.code.split()[-1] not in {"EA", "EB"} and urlsplit(page.final_url).hostname == urlsplit(source.website).hostname:
                from bs4 import BeautifulSoup
                identity = BeautifulSoup(page.text or "", "html.parser").find_all("meta", attrs={"name": lambda value: value and value.lower() in {"sitename", "publisher", "dc.publisher", "site_name"}})
                identity = [node for node in identity if node.get("content", "").strip()]
                if not identity:
                    if source.code not in self.identity_cache:
                        home = fetch(source.website, self.timeout, self.retries)
                        if (not home.status or home.status >= 400) and source.website.startswith("http://"):
                            home = fetch(secure_alternative(source.website), self.timeout, self.retries)
                        self.identity_cache[source.code] = home
                    home = self.identity_cache[source.code]
                    if home.status == 200 and urlsplit(home.final_url).hostname == urlsplit(source.website).hostname and authority_matches(source, home.text, home.final_url):
                        agency_ok = True
                        authority_evidence = self._evidence(source, home.final_url, home.content, phase="entry_identity")
            entries = discover_entries(page.text, page.final_url, source.allowed_hosts)
            energy = energy_entries(source, page.text, page.final_url)
            if energy:
                entries = energy + [item for item in entries if item["url"] not in {entry["url"] for entry in energy}]
            if links and not homepage and agency_ok:
                # Verified policy entries use their own pagination and configured roots.
                # General navigation must not divert a dated list to unrelated portals.
                entries = []
            # A homepage may help discovery but cannot be the verified policy entry.
            if homepage and entries:
                links = []
            discovery.append({"url": url, "final_url": page.final_url, "http_status": page.status,
                              "method": method, "kind": kind, "authority_match": agency_ok,
                              "authority_evidence": authority_evidence,
                              "energy_portal_inherited": energy_portal_inherited,
                              "list_original_evidence": list_original_evidence,
                              "candidate_count": len(links), "dated_candidates": sum(bool(item.get("date")) for item in links), "homepage": homepage,
                              "error": page.error, "evidence_path": page_evidence, "api_attempts": api_attempts})
            for link in links:
                if since and link.get("date") and link["date"] < since.isoformat():
                    continue
                if until and link.get("date") and link["date"] > until.isoformat():
                    continue
                decision = energy_policy_candidate(link)
                if decision["decision"] != "Collect":
                    scope_filtered.setdefault(normalize_url(link["url"]), {
                        "agency": source.code, "title": link["text"], "detail_url": link["url"],
                        "issue_date": link.get("date", ""), "list_url": page.final_url,
                        "list_evidence_path": page_evidence, "body_fetched": False, **decision})
                    continue
                origin = list_origins.get(key, page.final_url) if '/module/web/jpage/dataproxy.jsp' in urlsplit(page.final_url).path else page.final_url
                candidates.setdefault(normalize_url(link["url"]), {**link, "list_url": origin,
                                       "entry_authority_match": agency_ok, "homepage_entry": homepage,
                                       "energy_portal_inherited": energy_portal_inherited,
                                       "scope_decision": decision})
            page_boundary = bool(since and links and all(item.get("date") and item["date"] < since.isoformat() for item in links))
            if page_boundary:
                boundary_seen = True
            if len(candidates) >= limit:
                limited = bool(since) and (len(candidates) > limit or not boundary_seen)
                break
            next_pages = pagination_links(page.text, page.final_url, source.allowed_hosts)
            if next_pages and not page_boundary and pagination_count < source.pagination_pages:
                for next_url in next_pages:
                    if normalize_url(next_url) not in visited:
                        list_origins[normalize_url(next_url)] = list_origins.get(key, page.final_url)
                        queue.append((next_url, depth, "pagination"))
                        pagination_count += 1
            elif next_pages and not page_boundary:
                limited = True
            if depth < 3:
                for entry in entries[:6]:
                    if normalize_url(entry["url"]) not in visited:
                        queue.append((entry["url"], depth + 1, "entry"))
        exhausted = not queue
        selected = sorted(candidates.values(), key=lambda item: item.get("date", ""), reverse=True)[:limit]
        rows = []
        for item in selected:
            if not item["entry_authority_match"]:
                row = {"agency":source.code,"agency_name":source.chinese_name or source.name,
                       "title":item['text'],"detail_url":item['url'],"list_url":item['list_url'],
                       "issue_date":item.get('date',''),"policy_number":"","body_selector":"",
                       "body_valid":False,"body_text":"","abstract":"","evidence_path":"",
                       "body_sha256":"","body_sources":[],"processing_evidence":[],"attachment_evidence":[],
                       "attempts":[],"http_status":"","collector_method":"not_fetched",
                       "review_status":"Failed","error":"SOURCE_AUTHORITY_MISMATCH","failure_stage":"entry_identity"}
            else:
                row = self._detail(source, item, item["list_url"])
            row["scope_decision"] = item["scope_decision"]
            rows.append(row)
        success_count = sum(row["review_status"] == "Collected" for row in rows)
        failed_count = len(rows) - success_count
        status = "Collected" if success_count and not failed_count else "Partial" if success_count else "Failed"
        reason = "" if success_count else "NO_VERIFIED_BODY" if rows else "NO_DETAIL_LINKS"
        if failed_count and success_count:
            reason = "DETAIL_FAILURES: " + "; ".join(sorted({row["error"] for row in rows if not row["body_valid"]}))[:500]
        coverage = "diagnostic_sample" if since is None else "incomplete" if limited or queue else "date_boundary_reached" if boundary_seen else "entry_pages_exhausted"
        entry_verified = any(not item["homepage"] and item["authority_match"] and item["candidate_count"] for item in discovery)
        if not rows and scope_filtered:
            status, reason = "Filtered", "NO_IN_SCOPE_CANDIDATES"
        if any(item["decision"] == "Needs review" for item in scope_filtered.values()):
            reason = (reason + "; ENERGY_TOPIC_UNCONFIRMED").strip("; ")
            if not rows:
                status = "ScopeReview"
            if status == "Collected":
                status = "Partial"
        if (not rows and not scope_filtered and since and source.body_acceptance == "Passed" and entry_verified and boundary_seen and coverage != "incomplete"
                and all(item["dated_candidates"] == item["candidate_count"] and item["http_status"] == 200 and not item["error"]
                        and not any(attempt.get("error") for attempt in item["api_attempts"]) for item in discovery)):
            status, reason = "NoNew", "VERIFIED_ENTRY_NO_CANDIDATES_IN_WINDOW"
        if coverage == "incomplete" and status == "Collected":
            status, reason = "Partial", "WEEKLY_COVERAGE_INCOMPLETE"
        return rows, {"agency": source.code, "status": status, "http_status": discovery[0]["http_status"] if discovery else "",
                      "list_url": source.list_url, "final_url": discovery[-1]["final_url"] if discovery else "",
                      "candidates": len(selected), "body_success": success_count, "body_failed": failed_count,
                      "pages_visited": len(visited), "pagination_pages": pagination_count,
                      "coverage": coverage, "entry_verified": entry_verified,
                      "error": reason, "discovery": discovery,
                      "scope_filtered_count": len(scope_filtered),
                      "scope_review_count": sum(item["decision"] == "Needs review" for item in scope_filtered.values()),
                      "scope_filtered": list(scope_filtered.values())}

    def collect(self, sources, max_items=50, since=None, until=None):
        rows, statuses = [], []
        for source in sources:
            try:
                details, status = self.collect_source(source, max_items, since=since, until=until) if since is not None else self.collect_source(source, max_items)
            except Exception as exc:
                details = []
                status = {"agency": source.code, "status": "Failed", "list_url": source.list_url,
                          "candidates": 0, "body_success": 0, "error": "COLLECTOR_ERROR: " + str(exc)[:200]}
            rows.extend(details)
            statuses.append(status)
        unique = {}
        for row in rows:
            key = normalize_url(row["detail_url"])
            if key not in unique or row["body_valid"]:
                unique[key] = row
        return list(unique.values()), statuses
