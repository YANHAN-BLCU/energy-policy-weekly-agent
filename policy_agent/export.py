"""MI-like weekly workbook and JSON export."""
from pathlib import Path
from datetime import datetime
import hashlib, json
from urllib.parse import urlsplit
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .paths import local_path
from .taxonomy import dictionary, export_labels

HEADERS=['Policy_ID', 'Policy_Name', 'Policy_Number', 'Issue_Date', 'Issuing_Agency', 'Issuing_Agency_Name', 'Supported_Industry', 'Policy_Tags', 'Policy_Abstract', 'Impact_Analysis', 'Full_Text_Link', 'Source_List_URL', 'Collected_At', 'Collector_Method', 'Collection_Status', 'AI_Status', 'Evidence_Path']

def policy_id(row):
    return "P-"+hashlib.sha256((row.get("detail_url","")+"|"+row.get("issue_date","")).encode()).hexdigest()[:16]

def export_workbook(records, statuses, output_path, run_meta):
    if any(not r.get("body_valid") or not r.get("body_text", "").strip() or r.get("review_status") == "Failed" for r in records):
        raise ValueError("Cannot export failed collection records as weekly policies")
    wb=Workbook(); ws=wb.active; ws.title="Weekly_Policies"
    header_fill=PatternFill("solid",fgColor="1F4E78"); header_font=Font(color="FFFFFF",bold=True)
    ws.append(HEADERS)
    for c in ws[1]: c.fill=header_fill; c.font=header_font; c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for r in records:
        labels = export_labels(r)
        detail_url = r.get("detail_url", "")
        if not isinstance(detail_url, str) or urlsplit(detail_url).scheme not in {"http", "https"} or not urlsplit(detail_url).netloc:
            raise ValueError("Full_Text_Link requires a valid HTTP(S) policy URL")
        ws.append([policy_id(r),r.get("title",""),r.get("policy_number",""),r.get("issue_date",""),labels["agency_codes"],labels["agency_names"],labels["industry"],labels["tags"],r.get("abstract",""),r.get("impact_analysis",""),r.get("detail_url",""),r.get("list_url",""),run_meta["collected_at"],r.get("collector_method",""),r.get("review_status",""),r.get("ai_status",""),r.get("evidence_path","" )])
        for column in (11, 12):
            cell = ws.cell(ws.max_row, column)
            if cell.value:
                if not str(cell.value).startswith(("https://", "http://")):
                    raise ValueError("Policy hyperlink must be an HTTP(S) URL")
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    widths=[18,42,24,14,24,58,30,38,64,64,52,52,22,24,18,22,58]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
    summary=wb.create_sheet("Run_Summary")
    summary.append(["Metric","Value"])
    for key, value in {
        "Run_ID": run_meta.get("run_id", ""),
        "Collected_At": run_meta.get("collected_at", ""),
        "Source_Count": run_meta.get("source_count", 0),
        "Policy_Count": len(records),
        "Verified_Body_Records": run_meta.get("verified_body_records", 0),
        "Failed_Detail_Records": run_meta.get("failed_detail_records", 0),
        "Date_Unverified_Records": run_meta.get("date_unverified_records", 0),
        "Scope_Filtered_Records": run_meta.get("scope_filtered_records", 0),
        "Scope_Review_Records": run_meta.get("scope_review_records", 0),
        "Failed_Or_Partial_Sources": run_meta.get("failed_sources", 0),
        "Host_AI_Analyzed": sum(r.get("ai_status") == "Host AI analyzed" for r in records),
        "Needs_Review": sum(r.get("review_status") != "Collected" for r in records),
        "Window_Start": run_meta.get("window_start", ""),
        "Window_End": run_meta.get("window_end", ""),
        "Delivery_Email": run_meta.get("delivery_email", ""),
        "Weekday": run_meta.get("weekday", ""),
        "Trigger_Time": run_meta.get("trigger_time", ""),
        "Delivery_Status_At_Export": run_meta.get("delivery_status", "pending"),
        "Delivery_Receipt": "delivery_receipt.json",
    }.items():
        summary.append([key, value])
    for c in summary[1]: c.fill=header_fill; c.font=header_font
    status_ws=wb.create_sheet("Source_Status"); status_ws.append(["Agency","Status","HTTP_Status","List_URL","Final_URL","Candidates","Body_Success","Body_Failed","Pages_Visited","Entry_Verified","Coverage","Pagination_Pages","Error"]);
    for s in statuses: status_ws.append([s.get("agency"),s.get("status"),s.get("http_status"),s.get("list_url"),s.get("final_url",""),s.get("candidates",0),s.get("body_success",0),s.get("body_failed",0),s.get("pages_visited",0),s.get("entry_verified",False),s.get("coverage",""),s.get("pagination_pages",0),s.get("error","")])
    tags=wb.create_sheet("Tag_Dictionary"); tags.append(["Category","Tag"]);
    for item in dictionary()["policy_tags"]:
        tags.append([item["category"], item["tag"]])
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines=False
        if sheet.max_column:
            for cell in sheet[1]: cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    output_path=local_path(output_path, "表单路径"); output_path.parent.mkdir(parents=True,exist_ok=True); wb.save(output_path)
    return str(output_path)

def export_json(records,statuses,path,meta):
    local_path(path, "导出 JSON 路径").write_text(json.dumps({"meta":meta,"records":records,"statuses":statuses},ensure_ascii=False,indent=2),encoding="utf-8")
